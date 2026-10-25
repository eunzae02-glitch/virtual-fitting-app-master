"""키·체중·성별 기반 신체 치수 모델. 옷 치수나 사진상의 폭으로 환산하지 않는다.

모델 선택: 학습 자료의 교차검증 MAE.
예측 범위: 별도 보정 자료의 절대 잔차로 split-conformal 구간 계산.
최종 평가: 선택/보정에 사용하지 않은 사람들. 구간은 성별·부위별 주변적 범위다.
"""
import hashlib
import math
from collections import Counter
from dataclasses import dataclass
from pathlib import Path

import numpy as np


FEATURES = ('height_cm', 'weight_kg')
TARGETS = {
    'shoulder_breadth_cm': '어깨사이너비',
    'body_length_cm': '신체 기준 총장',
    'chest_circumference_cm': '가슴둘레',
}
HEADERS = {
    'HUMAN_ID': 'human_id', '성별': 'sex', '키(cm)': 'height_cm',
    '몸무게(kg)': 'weight_kg', '총장(cm)': 'body_length_cm',
    '어깨사이너비(cm)': 'shoulder_breadth_cm', '가슴둘레(cm)': 'chest_circumference_cm',
}
MODEL_NAMES = {'mean': '성별 평균', 'ridge': 'Ridge', 'gradient_boosting': 'Gradient Boosting'}
SEX_NAMES = {'M': '남', 'F': '여'}
ROOT = Path(__file__).resolve().parent.parent
DEFAULT_DATA = ROOT / 'size_data' / '전처리데이터.xlsx'
DEFAULT_OUTPUT = ROOT / 'outputs' / 'body_model'


@dataclass
class BodyDataset:
    ids: np.ndarray
    sexes: np.ndarray
    x: np.ndarray
    y: dict
    audit: dict


def normalize_sex(value):
    aliases = {'M': 'M', '남': 'M', '남자': 'M', '남성': 'M',
               'F': 'F', '여': 'F', '여자': 'F', '여성': 'F'}
    sex = aliases.get(str(value).strip().upper())
    if sex is None:
        raise ValueError('성별을 남 또는 여로 선택해주세요.')
    return sex


def positive_number(value, label):
    try:
        number = float(value)
    except (ValueError, TypeError, OverflowError) as exc:
        raise ValueError(f'{label}에 숫자를 입력해주세요.') from exc
    if isinstance(value, bool) or not math.isfinite(number) or number <= 0:
        raise ValueError(f'{label}는 0보다 큰 유한한 숫자여야 합니다.')
    return number


def load_dataset(path=DEFAULT_DATA, sheet='전처리', header_row=2):
    """단위가 명시된 전처리 파일만 읽는다. 원본 엑셀은 수정하지 않는다."""
    from openpyxl import load_workbook

    path = Path(path)
    workbook = load_workbook(path, read_only=True, data_only=True)
    records, dropped = [], Counter()
    seen = set()
    try:
        if sheet not in workbook.sheetnames:
            raise ValueError(f'엑셀에 {sheet!r} 시트가 없습니다.')
        rows = workbook[sheet].iter_rows(min_row=header_row, values_only=True)
        header = next(rows, ())
        columns = {}
        for index, name in enumerate(header):
            key = HEADERS.get(''.join(str(name or '').split()))
            if key:
                if key in columns:
                    raise ValueError(f'엑셀 헤더가 중복됐습니다: {key}')
                columns[key] = index
        missing = set(HEADERS.values()) - set(columns)
        if missing:
            raise ValueError('엑셀 헤더·단위·헤더 행을 확인해주세요: ' + ', '.join(sorted(missing)))
        for row in rows:
            if not any(value is not None for value in row):
                continue
            raw = {key: row[index] if index < len(row) else None
                   for key, index in columns.items()}
            human_id = str(raw['human_id'] or '').strip()
            if not human_id:
                dropped['HUMAN_ID 누락'] += 1
                continue
            if human_id in seen:
                raise ValueError('중복 HUMAN_ID가 있습니다. 사람별 중복 자료를 먼저 확인해주세요.')
            seen.add(human_id)
            try:
                record = {'human_id': human_id, 'sex': normalize_sex(raw['sex'])}
                for key in FEATURES + tuple(TARGETS):
                    record[key] = positive_number(raw[key], key)
                # 단위 오류/특수 코드를 정상 실측으로 학습하지 않는다. 이상치는 자동 보정하지 않는다.
                if any(record[key] in (999, 9999) for key in FEATURES + tuple(TARGETS)):
                    raise ValueError('결측·특수 코드 확인 필요')
                if record['height_cm'] >= 300:
                    raise ValueError('키의 cm 단위 확인 필요')
            except ValueError as exc:
                dropped[str(exc)] += 1
                continue
            records.append(record)
    finally:
        workbook.close()
    if not records:
        raise ValueError('학습할 유효 데이터가 없습니다.')
    records.sort(key=lambda row: row['human_id'])
    audit = {
        'source_file': path.name, 'source_sha256': hashlib.sha256(path.read_bytes()).hexdigest(),
        'sheet': sheet, 'header_row': header_row, 'accepted_rows': len(records),
        'dropped_rows': sum(dropped.values()), 'drop_reasons': dict(dropped),
        'sex_counts': dict(Counter(row['sex'] for row in records)),
        'units': {'height': 'cm', 'weight': 'kg', 'targets': 'cm'},
        'body_length_definition': '목뒤높이 - 위앞엉덩뼈가시높이',
    }
    return BodyDataset(
        np.array([row['human_id'] for row in records]),
        np.array([row['sex'] for row in records]),
        np.array([[row[key] for key in FEATURES] for row in records], dtype=float),
        {key: np.array([row[key] for row in records], dtype=float) for key in TARGETS}, audit,
    )


def split_people(dataset, seed=42):
    """사람별 고정 60/20/20 분할. 성별 비율을 유지한다."""
    from sklearn.model_selection import train_test_split

    if any(np.sum(dataset.sexes == sex) < 30 for sex in SEX_NAMES):
        raise ValueError('남녀 각각 30명 이상의 유효 데이터가 필요합니다.')
    indices = np.arange(len(dataset.ids))
    remaining, test = train_test_split(indices, test_size=0.2, random_state=seed,
                                       stratify=dataset.sexes)
    train, calibration = train_test_split(remaining, test_size=0.25, random_state=seed,
                                          stratify=dataset.sexes[remaining])
    return {'train': train, 'calibration': calibration, 'test': test}


def conformal_radius(actual, predicted, alpha=0.1):
    """유한 표본 보정을 적용한 절대 잔차 분위수. 최종 평가 자료는 사용하지 않는다."""
    if not 0 < alpha < 1:
        raise ValueError('alpha는 0과 1 사이여야 합니다.')
    errors = np.sort(np.abs(np.asarray(actual, dtype=float) - np.asarray(predicted, dtype=float)))
    if errors.ndim != 1 or not len(errors) or not np.all(np.isfinite(errors)):
        raise ValueError('예측 범위 보정 자료를 확인해주세요.')
    rank = math.ceil((len(errors) + 1) * (1 - alpha))
    if rank > len(errors):
        raise ValueError('선택한 예측 범위를 계산하기에 보정 자료가 부족합니다.')
    return float(errors[rank - 1])


def fit_candidates(x, y, seed=42, folds=5):
    from sklearn.dummy import DummyRegressor
    from sklearn.ensemble import GradientBoostingRegressor
    from sklearn.linear_model import Ridge
    from sklearn.model_selection import GridSearchCV, KFold, cross_val_score
    from sklearn.pipeline import make_pipeline
    from sklearn.preprocessing import StandardScaler

    cv = list(KFold(n_splits=folds, shuffle=True, random_state=seed).split(x))
    mean = DummyRegressor(strategy='mean')
    mean_scores = -cross_val_score(mean, x, y, cv=cv, scoring='neg_mean_absolute_error')
    candidates = {
        'mean': {'estimator': mean.fit(x, y), 'cv_mae_cm': float(mean_scores.mean()),
                 'cv_mae_std_cm': float(mean_scores.std()), 'parameters': {}},
    }
    searches = {
        'ridge': (make_pipeline(StandardScaler(), Ridge()),
                  {'ridge__alpha': [0.1, 1.0, 10.0, 100.0]}),
        'gradient_boosting': (
            GradientBoostingRegressor(random_state=seed, learning_rate=0.05, min_samples_leaf=10),
            {'n_estimators': [100, 200], 'max_depth': [1, 2]},
        ),
    }
    for name, (estimator, parameters) in searches.items():
        search = GridSearchCV(estimator, parameters, scoring='neg_mean_absolute_error',
                              cv=cv, n_jobs=1, error_score='raise')
        search.fit(x, y)
        candidates[name] = {
            'estimator': search.best_estimator_, 'cv_mae_cm': float(-search.best_score_),
            'cv_mae_std_cm': float(search.cv_results_['std_test_score'][search.best_index_]),
            'parameters': search.best_params_,
        }
    # 동률이면 평균→Ridge→Boosting 순. 최종 평가 결과로 모델을 다시 고르지 않는다.
    selected = min(candidates, key=lambda name: candidates[name]['cv_mae_cm'])
    return candidates, selected


def train_models(dataset, seed=42, folds=5, alpha=0.1, progress=None):
    from importlib.metadata import version

    splits = split_people(dataset, seed)
    report = {
        'schema_version': 1, 'dataset': dataset.audit, 'seed': seed, 'folds': folds,
        'coverage_target': 1 - alpha,
        'selection_rule': '성별·항목별 학습 자료 교차검증 MAE 최저. 최종 평가·범위 보정 자료로 선택하지 않음.',
        'interval_method': '성별·부위별 split-conformal 절대 잔차. 개인별·세 부위 동시 포함 보장 아님.',
        'versions': {name: version(name) for name in ('numpy', 'scikit-learn', 'openpyxl', 'joblib')},
        'split_counts': {name: {sex: int(np.sum(dataset.sexes[index] == sex))
                                for sex in SEX_NAMES} for name, index in splits.items()},
        'metrics': [],
    }
    bundle = {'schema_version': 1, 'feature_names': FEATURES, 'models': {},
              'feature_ranges': {}, 'report': report}
    predictions = []
    for sex in SEX_NAMES:
        groups = {name: index[dataset.sexes[index] == sex] for name, index in splits.items()}
        train, cal, test = (groups[name] for name in ('train', 'calibration', 'test'))
        bundle['models'][sex] = {}
        bundle['feature_ranges'][sex] = {
            key: [float(dataset.x[train, column].min()), float(dataset.x[train, column].max())]
            for column, key in enumerate(FEATURES)
        }
        for target, label in TARGETS.items():
            if progress:
                progress(f'{SEX_NAMES[sex]} / {label}: 평균·Ridge·Gradient Boosting 학습 중')
            y = dataset.y[target]
            candidates, selected = fit_candidates(dataset.x[train], y[train], seed, folds)
            for name, candidate in candidates.items():
                estimator = candidate['estimator']
                radius = conformal_radius(y[cal], estimator.predict(dataset.x[cal]), alpha)
                candidate['interval_radius_cm'] = radius
                predicted = estimator.predict(dataset.x[test])
                errors = predicted - y[test]
                metric = {
                    'sex': sex, 'target': target, 'model': name, 'selected': name == selected,
                    'cv_mae_cm': candidate['cv_mae_cm'], 'cv_mae_std_cm': candidate['cv_mae_std_cm'],
                    'test_mae_cm': float(np.mean(np.abs(errors))),
                    'test_rmse_cm': float(np.sqrt(np.mean(errors ** 2))),
                    'test_bias_cm': float(np.mean(errors)),
                    'test_p90_absolute_error_cm': float(np.quantile(np.abs(errors), 0.9)),
                    'interval_radius_cm': radius,
                    'test_coverage': float(np.mean(np.abs(errors) <= radius)),
                    'parameters': candidate['parameters'],
                }
                report['metrics'].append(metric)
                for index, value in zip(test, predicted):
                    predictions.append({
                        'human_id': str(dataset.ids[index]), 'sex': sex, 'target': target, 'model': name,
                        'actual_cm': float(y[index]), 'predicted_cm': float(value),
                        'error_cm': float(value - y[index]),
                        'lower_cm': float(value - radius), 'upper_cm': float(value + radius),
                    })
            bundle['models'][sex][target] = {'selected': selected, 'candidates': candidates}
    return bundle, splits, predictions


def report_markdown(report):
    audit = report['dataset']
    lines = ['# 성별·키·체중 기반 신체 치수 모델 비교', '',
             f"자료: {audit['source_file']} / 유효 {audit['accepted_rows']}명 / 제외 {audit['dropped_rows']}명.",
             f"SHA256: `{audit['source_sha256']}`", '',
             '입력: 키(cm)·체중(kg). 성별로 모델을 분리한다. ID와 정답 치수는 입력 특성에서 제외한다.',
             '신체 기준 총장 = 목뒤높이 − 위앞엉덩뼈가시높이. 옷 총장이나 사진 픽셀 폭과 구분한다.', '',
             f"선택 기준: {report['selection_rule']}",
             f"예측 구간: {report['interval_method']}", '',
             '| 성별 | 학습 | 범위 보정 | 최종 평가 |', '|---|---:|---:|---:|']
    for sex in SEX_NAMES:
        counts = [report['split_counts'][name][sex] for name in ('train', 'calibration', 'test')]
        lines.append(f'| {SEX_NAMES[sex]} | {counts[0]} | {counts[1]} | {counts[2]} |')
    lines += ['', '| 성별 | 예측 항목 | 모델 | CV MAE(cm) | 평가 MAE(cm) | RMSE(cm) | 구간 ±cm | 평가 포함률 | 선택 |',
              '|---|---|---|---:|---:|---:|---:|---:|---|']
    for row in report['metrics']:
        lines.append(f"| {SEX_NAMES[row['sex']]} | {TARGETS[row['target']]} | {MODEL_NAMES[row['model']]} | "
                     f"{row['cv_mae_cm']:.3f} | {row['test_mae_cm']:.3f} | {row['test_rmse_cm']:.3f} | "
                     f"{row['interval_radius_cm']:.3f} | {row['test_coverage']:.1%} | "
                     f"{'선택' if row['selected'] else ''} |")
    lines += ['', f"구간 목표 포함률: {report['coverage_target']:.0%}. 실제 포함률은 위 평가 결과를 따른다.",
              '최종 평가 MAE가 더 작은 다른 모델이 있어도 이 결과를 보고 모델을 다시 선택하지 않는다.',
              '같은 조사자료 내 무작위 분할 평가이며, 외부 사용자·다른 조사집단에서의 정확도 검증은 별도다.',
              '예측값은 몸 치수다. 어깨너비·총장·가슴단면의 옷 실측으로 대체하지 않는다.', '',
              f"실행 설정: seed={report['seed']}, CV={report['folds']} folds.",
              f"라이브러리: {report['versions']}"]
    return '\n'.join(lines) + '\n'


def latest_model_path():
    paths = list(DEFAULT_OUTPUT.glob('*/model.joblib'))
    if not paths:
        raise FileNotFoundError('학습 모델이 없습니다. 먼저 scripts/train_body_model.py를 실행해주세요.')
    return max(paths, key=lambda path: path.stat().st_mtime_ns)


def load_model(path=None):
    import joblib

    bundle = joblib.load(Path(path) if path else latest_model_path())
    if bundle.get('schema_version') != 1 or tuple(bundle.get('feature_names', ())) != FEATURES:
        raise ValueError('지원하지 않는 신체 예측 모델 형식입니다.')
    return bundle


def predict_body(bundle, sex, height_cm, weight_kg, model='selected'):
    sex = normalize_sex(sex)
    values = [positive_number(height_cm, '키(cm)'), positive_number(weight_kg, '체중(kg)')]
    for key, label, value in zip(FEATURES, ('키', '체중'), values):
        low, high = bundle['feature_ranges'][sex][key]
        if not low <= value <= high:
            unit = 'cm' if key == 'height_cm' else 'kg'
            raise ValueError(f'{label}가 {SEX_NAMES[sex]} 학습 자료 범위({low:g}~{high:g}{unit}) 밖입니다. '
                             '실측 또는 기준 옷 정보를 사용해주세요.')
    if model != 'selected' and model not in MODEL_NAMES:
        raise ValueError('비교할 모델을 확인해주세요.')
    output = {}
    for key, label in TARGETS.items():
        group = bundle['models'][sex][key]
        name = group['selected'] if model == 'selected' else model
        candidate = group['candidates'][name]
        value = float(candidate['estimator'].predict(np.asarray([values], dtype=float))[0])
        radius = candidate['interval_radius_cm']
        output[key] = {'label': label, 'value_cm': value, 'low_cm': value - radius,
                       'high_cm': value + radius, 'model': name, 'source': 'excel_body_model'}
    return output
