"""저장된 신체 모델 평가 결과로 현재 진행 현황 발표 자료를 만든다.

실행: python -X utf8 scripts/build_body_model_progress_ppt.py
필요 패키지: python-pptx
"""

import json
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / 'docs' / 'fitcheck_body_model_progress_20261025.pptx'
RUN = ROOT / 'outputs' / 'body_model' / '20261025T051648_730977Z'
REPORT = RUN / 'report.json'

NAVY = RGBColor(14, 26, 46)
INK = RGBColor(31, 47, 66)
MUTED = RGBColor(95, 110, 126)
LIGHT = RGBColor(239, 244, 247)
WHITE = RGBColor(255, 255, 255)
TEAL = RGBColor(0, 156, 149)
PALE_TEAL = RGBColor(222, 245, 242)
CORAL = RGBColor(237, 111, 92)
PALE_CORAL = RGBColor(255, 237, 231)
GRAY = RGBColor(216, 224, 230)
FONT = '맑은 고딕'


def rect(slide, x, y, w, h, color, radius=False):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    if radius:
        shape.adjustments[0] = 0.10
    return shape


def label(slide, value, x, y, w, h, size=18, color=INK, bold=False,
          align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.MIDDLE):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = 0
    frame.margin_right = 0
    frame.margin_top = 0
    frame.margin_bottom = 0
    frame.vertical_anchor = valign
    for index, line in enumerate(value.split('\n')):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.alignment = align
        paragraph.space_after = Pt(1)
        run = paragraph.add_run()
        run.text = line
        run.font.name = FONT
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
    return box


def base_slide(prs, number, eyebrow, title, subtitle=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = WHITE
    rect(slide, 0, 0, 13.333, .11, TEAL)
    label(slide, eyebrow, .62, .32, 8, .30, 11, TEAL, True)
    label(slide, title, .62, .76, 12, .65, 28, NAVY, True)
    if subtitle:
        label(slide, subtitle, .63, 1.44, 12.1, .45, 13, MUTED)
    rect(slide, .62, 7.05, 12.08, .012, GRAY)
    label(slide, 'FitCheck · 신체 치수 모델 진행 현황 · 2026.10.25', .62, 7.10, 9, .24, 10, MUTED)
    label(slide, f'{number:02d}', 12.12, 7.10, .58, .24, 10, MUTED, align=PP_ALIGN.RIGHT)
    return slide


def card(slide, x, y, w, h, headline, body, accent=TEAL, fill=LIGHT,
         headline_size=17, body_size=13):
    rect(slide, x, y, w, h, fill, radius=True)
    rect(slide, x, y, .075, h, accent)
    label(slide, headline, x + .22, y + .17, w - .42, .45, headline_size, NAVY, True)
    label(slide, body, x + .22, y + .68, w - .42, h - .84, body_size, INK,
          valign=MSO_ANCHOR.TOP)


def table_row(slide, values, y, widths, x=.80, h=.52, fill=None,
              header=False, highlight=None):
    if fill:
        rect(slide, x, y, sum(widths), h, fill)
    cursor = x
    for i, (value, width) in enumerate(zip(values, widths)):
        color = TEAL if highlight == i else (MUTED if header else INK)
        label(slide, str(value), cursor + .10, y + .04, width - .18, h - .08,
              11 if header else 13, color, header or highlight == i,
              align=PP_ALIGN.LEFT if i < 2 else PP_ALIGN.CENTER)
        cursor += width


def metric(metrics, sex, target, model):
    return next(item for item in metrics if item['sex'] == sex
                and item['target'] == target and item['model'] == model)


def create_deck(report):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    metrics = report['metrics']
    target_names = {
        'shoulder_breadth_cm': '어깨사이너비',
        'body_length_cm': '신체 기준 총장',
        'chest_circumference_cm': '가슴둘레',
    }

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = NAVY
    rect(slide, .74, .88, .11, 5.61, TEAL)
    label(slide, 'FITCHECK  /  CAPSTONE PROJECT', 1.15, 1.00, 10.5, .45, 16, TEAL, True)
    label(slide, '키·체중·성별 기반\n신체 치수 예측', 1.12, 1.78, 11.2, 1.72,
          37, WHITE, True, valign=MSO_ANCHOR.TOP)
    label(slide, '성별 평균 · Ridge · Gradient Boosting 비교 결과', 1.15, 4.12,
          10.6, .48, 20, WHITE)
    rect(slide, 1.12, 5.23, 4.55, .54, TEAL, radius=True)
    label(slide, '최종 선택  RIDGE 회귀', 1.30, 5.33, 4.14, .31, 15, WHITE, True)
    label(slide, '현재 진행 현황  |  2026.10.25', 1.15, 6.60, 8, .36, 12, RGBColor(168, 190, 202))

    slide = base_slide(prs, 2, '01 / 프로젝트 방향', '기준 옷이 있으면 옷끼리, 없으면 신체 예측으로',
                       '상의 사이즈 비교와 향후 합성 조정을 위한 두 가지 입력 경로')
    card(slide, .75, 2.10, 5.65, 2.65, '경로 A  |  잘 맞는 기준 상의',
         '기준 옷의 총장 · 어깨너비 · 가슴단면을 직접 입력\n→ 새 옷 판매 사이즈와 항목별 비교', TEAL, PALE_TEAL)
    card(slide, 6.91, 2.10, 5.65, 2.65, '경로 B  |  기준 옷이 없는 경우',
         '성별 · 키 · 체중 입력 → 신체 치수 범위 예측\n→ 새 옷 실측과 여유분 비교', CORAL, PALE_CORAL)
    rect(slide, .75, 5.20, 11.81, .94, LIGHT, radius=True)
    label(slide, '이번 발표 범위', 1.02, 5.37, 2.1, .48, 15, NAVY, True)
    label(slide, '기준 옷 직접 입력 + 엑셀 모델 학습·비교·웹 예측까지 완료',
          3.09, 5.37, 9.1, .48, 16, TEAL, True)

    slide = base_slide(prs, 3, '02 / 데이터와 예측값', '실측 데이터 5,092명 · 예측 대상은 3개',
                       'size_data/전처리데이터.xlsx  ·  남 2,319명 / 여 2,773명  ·  제외 0명')
    card(slide, .75, 2.13, 3.74, 2.31, '입력', '사용자가 선택한 성별\n키(cm) · 체중(kg)', TEAL, LIGHT)
    card(slide, 4.80, 2.13, 3.74, 2.31, '예측', '어깨사이너비(cm)\n신체 기준 총장(cm)\n가슴둘레(cm)', TEAL, LIGHT)
    card(slide, 8.85, 2.13, 3.74, 2.31, '옷 치수와 비교',
         '옷 어깨너비 · 옷 총장\n옷 가슴단면', TEAL, LIGHT)
    rect(slide, .75, 4.80, 11.84, 1.33, PALE_CORAL, radius=True)
    label(slide, '측정 정의', 1.00, 5.05, 1.50, .44, 15, CORAL, True)
    label(slide, '신체 기준 총장 = 목뒤높이 − 위앞엉덩뼈가시높이\n몸 가슴둘레 / 2는 옷 가슴단면과 비교할 때만 사용',
          2.62, 5.00, 9.43, .89, 15, INK)

    slide = base_slide(prs, 4, '03 / 평가 설계', '모델 선택과 최종 평가는 사람 단위로 분리',
                       '학습 자료의 교차검증 MAE만 보고 모델을 선택')
    for x, count, title, detail in [
        (.76, '3,054명', '학습 60%', '5-fold CV로 평균·Ridge·Boosting 비교'),
        (4.94, '1,019명', '범위 보정 20%', '절대 오차로 90% 목표 예측 구간 계산'),
        (9.12, '1,019명', '최종 평가 20%', 'MAE와 구간 실제 포함률 확인'),
    ]:
        rect(slide, x, 2.27, 3.42, 2.24, LIGHT, radius=True)
        label(slide, count, x + .20, 2.50, 3.0, .58, 27, TEAL, True)
        label(slide, title, x + .20, 3.12, 3.0, .44, 18, NAVY, True)
        label(slide, detail, x + .20, 3.69, 3.0, .62, 12, INK,
              valign=MSO_ANCHOR.TOP)
    label(slide, '→', 4.32, 3.06, .53, .50, 25, TEAL, True, PP_ALIGN.CENTER)
    label(slide, '→', 8.50, 3.06, .53, .50, 25, TEAL, True, PP_ALIGN.CENTER)
    rect(slide, .76, 5.02, 11.78, 1.11, PALE_TEAL, radius=True)
    label(slide, '누출 방지', 1.02, 5.30, 1.70, .42, 16, TEAL, True)
    label(slide, 'ID·정답 치수를 입력 특성에서 제외 · 사람 중복 없는 분할 · 최종 평가로 모델 재선택 없음',
          2.55, 5.30, 9.50, .44, 14, INK)

    slide = base_slide(prs, 5, '04 / 모델 선택', '6개 성별·항목 조합 모두 Ridge가 CV 최저',
                       '단위: cm · 낮을수록 좋음 · 아래 수치는 학습 자료 5-fold 교차검증 MAE')
    widths = [1.09, 2.76, 2.18, 2.18, 2.18]
    table_row(slide, ['성별', '예측 항목', '성별 평균', 'Ridge', 'Boosting'],
              2.04, widths, h=.53, fill=PALE_TEAL, header=True)
    y = 2.58
    for sex, sex_label in [('M', '남'), ('F', '여')]:
        for target, target_name in target_names.items():
            values = [sex_label, target_name]
            values.extend(f"{metric(metrics, sex, target, model)['cv_mae_cm']:.3f}"
                          for model in ['mean', 'ridge', 'gradient_boosting'])
            table_row(slide, values, y, widths, h=.53,
                      fill=WHITE if (int((y - 2.58) / .53) % 2 == 0) else LIGHT,
                      highlight=3)
            y += .53
    rect(slide, .80, 6.06, 11.37, .51, PALE_CORAL, radius=True)
    label(slide, 'Ridge와 Boosting 차이는 일부 항목에서 작다. 선택 기준은 CV이며 최종 평가값으로 바꾸지 않았다.',
          1.02, 6.17, 10.95, .28, 12, INK)

    slide = base_slide(prs, 6, '05 / 최종 평가', '선택된 Ridge의 성별·항목별 오차',
                       '최종 평가 1,019명 · 예측 구간은 별도 보정 자료로 만든 90% 목표 범위')
    widths = [1.05, 3.30, 2.15, 2.15, 2.50]
    table_row(slide, ['성별', '예측 항목', '평가 MAE', '구간 ±', '실제 포함률'],
              2.02, widths, h=.53, fill=PALE_TEAL, header=True)
    y = 2.55
    for sex, sex_label in [('M', '남'), ('F', '여')]:
        for target, target_name in target_names.items():
            row = metric(metrics, sex, target, 'ridge')
            table_row(slide, [sex_label, target_name,
                              f"{row['test_mae_cm']:.2f} cm",
                              f"{row['interval_radius_cm']:.2f} cm",
                              f"{row['test_coverage']:.1%}"],
                      y, widths, h=.53, fill=WHITE if y < 4.14 else LIGHT,
                      highlight=2)
            y += .53
    rect(slide, .80, 6.04, 11.15, .53, PALE_CORAL, radius=True)
    label(slide, '예측 구간은 성별·부위별 주변적 범위이며 한 사람의 세 치수를 동시에 포함한다는 보장은 없다.',
          1.04, 6.16, 10.78, .28, 12, INK)

    slide = base_slide(prs, 7, '06 / 동작 확인', '현재 구현된 웹 디버깅 화면은 2개',
                       '직접 입력 경로와 모델 예측 경로를 각각 검증할 수 있다')
    card(slide, .78, 2.13, 5.68, 3.18, '1단계  |  기준 옷 치수 입력',
         '총장 · 어깨너비 · 가슴단면 입력 및 검증\n값 수정 시 확인 상태 초기화\n실행: scripts/debug_reference_size.py --ui\n주소: localhost:7863',
         TEAL, PALE_TEAL, 18, 13)
    card(slide, 6.82, 2.13, 5.68, 3.18, '2단계  |  신체 모델 예측',
         '성별 · 키 · 체중 → 세 신체 치수와 범위\n모델별 결과 / 선택적 실측 오차 대조\n실행: scripts/debug_body_model.py --ui\n주소: localhost:7864',
         CORAL, PALE_CORAL, 18, 13)
    label(slide, '검증: 관련 테스트 9개 통과 · 저장 모델 불러오기 및 웹 API 예측 확인',
          1.02, 5.79, 11.4, .41, 16, NAVY, True)

    slide = base_slide(prs, 8, '07 / 이전 실험', '사진 측정과 가상 피팅의 실험 결과를 보존',
                       '새 신체 예측 경로를 택하게 된 배경과 기존 합성 앱의 상태')
    card(slide, .79, 2.12, 5.68, 2.63, '사진 기반 치수 측정',
         'MediaPipe 어깨 관절점과 입력 키를 이용한 환산 실험\n사진 구도·원근·자세에 따라 결과가 흔들림\n고정 보정계수 1.28의 일반화는 미검증',
         CORAL, PALE_CORAL, 17, 13)
    card(slide, 6.81, 2.12, 5.68, 2.63, '기존 가상 피팅 앱',
         'CatVTON 속도·품질 실험 후 FASHN을 앱에 연결\nFitDiT도 비교했으나 정핏 조정 실험에서 한계 확인\n기존 합성은 새 치수 기반 조정과 별도',
         TEAL, PALE_TEAL, 17, 13)
    rect(slide, .79, 5.18, 11.70, .88, LIGHT, radius=True)
    label(slide, '판단', 1.03, 5.38, 1.07, .41, 16, TEAL, True)
    label(slide, '기준 옷 비교와 엑셀 예측을 신체 치수 입력 경로로 정했다.',
          2.09, 5.38, 9.95, .41, 17, INK, True)

    slide = base_slide(prs, 9, '08 / 해석과 한계', '오차 수치는 모델 성능이지 실제 착용 핏 정확도가 아니다',
                       '발표 시 신체 치수, 옷 실측, 사진상 폭을 구분해서 설명')
    card(slide, .79, 2.14, 5.68, 1.57, '사진 측정의 한계',
         '구도·원근·옷 때문에 어깨 픽셀 측정이 흔들려 기본 예측 경로에서 제외',
         CORAL, PALE_CORAL, 16, 12)
    card(slide, 6.81, 2.14, 5.68, 1.57, '평가의 한계',
         '동일 조사자료 내 무작위 분할 결과. 외부 사용자에서의 정확도는 미검증',
         CORAL, PALE_CORAL, 16, 12)
    card(slide, .79, 4.05, 5.68, 1.57, '측정 정의',
         '신체 기준 총장은 옷 총장과 다르고 몸 둘레의 절반은 사진 정면 폭이 아님',
         TEAL, LIGHT, 16, 12)
    card(slide, 6.81, 4.05, 5.68, 1.57, '현재 제품 범위',
         '새 옷 OCR·S/M/L 연결, 치수 기반 마스킹·합성 조정은 후속 작업',
         TEAL, LIGHT, 16, 12)

    slide = base_slide(prs, 10, '09 / 다음 단계', 'Ridge를 기준 모델로 두고 연결과 실측 검증 진행',
                       '새 옷 추천과 이미지 합성 조정은 각각 별도의 검증 단계를 거친다')
    steps = [
        ('01', '입력 경로 연결', '기준 옷이 있으면 옷 실측, 없으면 Ridge 신체 예측 사용'),
        ('02', '대상 옷 치수표', 'OCR 원문 확인·수정 후 S/M/L별 총장·어깨·가슴 비교'),
        ('03', '실측·외부 검증', '실제 사용자 오차, 체형별 실패, 추천 경계 사례 기록'),
        ('04', '합성 조정 실험', '어깨 마스크·세로 길이·가슴 영역 각각 효과 검증'),
    ]
    for i, (num, title, detail) in enumerate(steps):
        y = 2.08 + i * 1.09
        rect(slide, .79, y, 11.72, .88, LIGHT, radius=True)
        rect(slide, .94, y + .15, .57, .57, TEAL if i < 2 else CORAL, radius=True)
        label(slide, num, 1.03, y + .26, .38, .32, 12, WHITE, True, PP_ALIGN.CENTER)
        label(slide, title, 1.73, y + .22, 2.50, .39, 17, NAVY, True)
        label(slide, detail, 4.28, y + .22, 7.95, .41, 14, INK)

    slide = base_slide(prs, 11, '10 / 재현 자료', '학습 결과와 발표 수치를 다시 확인하는 방법',
                       '프로젝트 루트에서 실행 · 저장된 실행 폴더는 덮어쓰지 않음')
    rect(slide, .80, 2.11, 11.71, 2.22, NAVY, radius=True)
    commands = ('python -X utf8 scripts/train_body_model.py\n'
                'python -X utf8 scripts/debug_body_model.py --ui\n'
                'python -X utf8 scripts/debug_reference_size.py --ui')
    label(slide, commands, 1.14, 2.46, 10.95, 1.54, 19, WHITE,
          valign=MSO_ANCHOR.TOP)
    card(slide, .80, 4.65, 11.71, 1.38, '근거 파일',
         'outputs/body_model/20261025T051648_730977Z/report.md · report.json · metrics.csv',
         TEAL, PALE_TEAL, 16, 12)
    return prs


def main():
    with REPORT.open(encoding='utf-8') as source:
        report = json.load(source)
    prs = create_deck(report)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUTPUT)
    print(f'{len(prs.slides)}장 저장: {OUTPUT}')


if __name__ == '__main__':
    main()
