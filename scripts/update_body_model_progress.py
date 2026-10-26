"""기존 신체 모델 발표 자료에 이후 사이즈 합성 진행 상황을 덧붙인다."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / 'docs' / 'fitcheck_body_model_progress_20261025.pptx'
OUTPUT = ROOT / 'docs' / 'fitcheck_progress_20261026.pptx'

INK = RGBColor(14, 26, 46)
MUTED = RGBColor(95, 110, 126)
TEAL = RGBColor(0, 156, 149)
PALE = RGBColor(237, 248, 247)
SURFACE = RGBColor(246, 248, 250)
LINE = RGBColor(221, 229, 234)
ORANGE = RGBColor(199, 107, 31)
FONT = '맑은 고딕'


def text(slide, value, x, y, w, h, size=16, color=INK, bold=False, align=None):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(0.02)
    frame.margin_right = Inches(0.02)
    frame.margin_top = 0
    frame.margin_bottom = 0
    for index, line in enumerate(value.split('\n')):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        if align is not None:
            paragraph.alignment = align
        paragraph.space_after = Pt(3)
        run = paragraph.add_run()
        run.text = line
        run.font.name = FONT
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
    return shape


def box(slide, x, y, w, h, fill=SURFACE, stroke=LINE, rounded=True):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = stroke
    return shape


def header(presentation, section, title, subtitle):
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor(255, 255, 255)
    box(slide, 0, 0, 13.333, .11, TEAL, TEAL, rounded=False)
    text(slide, section, .62, .32, 8, .3, 11, TEAL, True)
    text(slide, title, .62, .76, 12, .65, 28, INK, True)
    text(slide, subtitle, .63, 1.44, 12.1, .45, 13, MUTED)
    box(slide, .62, 7.05, 12.08, .012, LINE, LINE, rounded=False)
    text(slide, 'FitCheck · 사이즈 합성 진행 현황 · 2026.10.26', .62, 7.1, 9, .24, 10, MUTED)
    text(slide, f'{len(presentation.slides):02d}', 12.12, 7.1, .58, .24, 10, MUTED)
    return slide


def card(slide, x, y, w, h, kicker, title, lines, accent=TEAL):
    box(slide, x, y, w, h)
    box(slide, x, y, .075, h, accent, accent, rounded=False)
    text(slide, kicker, x+.23, y+.2, w-.45, .28, 11, accent, True)
    text(slide, title, x+.23, y+.6, w-.45, .52, 20, INK, True)
    for index, line in enumerate(lines):
        text(slide, line, x+.23, y+1.34+index*.57, w-.48, .46, 13, MUTED)


def replace_once(slide, before, after):
    for shape in slide.shapes:
        if shape.has_text_frame and before in shape.text:
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    if before in run.text:
                        run.text = run.text.replace(before, after)
                        return
    raise ValueError(f'원본 슬라이드 문구를 찾지 못했습니다: {before}')


def update_old_status(presentation):
    # 기존 학습·평가 기록은 보존하고, 이후 구현으로 오래된 상태 문구만 갱신한다.
    replace_once(presentation.slides[0], '신체 치수 예측', '사이즈별 합성 실험까지')
    replace_once(presentation.slides[0], '키·체중·성별 기반', '신체 치수 예측부터')
    replace_once(presentation.slides[0], '성별 평균 · Ridge · Gradient Boosting 비교 결과',
                 'Ridge 예측 · 판매 사이즈 비교 · FASHN 후처리')
    replace_once(presentation.slides[0], '현재 진행 현황  |  2026.10.25',
                 '현재 진행 현황  |  2026.10.26')
    replace_once(presentation.slides[0], '최종 선택  RIDGE 회귀',
                 '신체 예측 모델  RIDGE 회귀')
    replace_once(presentation.slides[1], '이번 발표 범위', '초기 구현 단계')
    replace_once(presentation.slides[6], '현재 구현된 웹 디버깅 화면은 2개',
                 '첫 단계에서 만든 신체 치수 디버거 2개')
    replace_once(presentation.slides[7], '기존 합성은 새 치수 기반 조정과 별도',
                 '이후 FASHN 기본 합성에 치수 후처리를 연결')
    replace_once(presentation.slides[8], '새 옷 OCR·S/M/L 연결, 치수 기반 마스킹·합성 조정은 후속 작업',
                 'S/M/L 치수 비교·후처리 실험 연결; OCR·정확도 검증은 후속')
    replace_once(presentation.slides[9], '09 / 다음 단계', '09 / 이후 진행 방향')
    replace_once(presentation.slides[9], 'Ridge를 기준 모델로 두고 연결과 실측 검증 진행',
                 'Ridge 비교에서 FASHN 합성 후처리까지 확장')
    replace_once(presentation.slides[9], '입력 경로 연결', '입력 경로 연결 완료')
    replace_once(presentation.slides[9], '대상 옷 치수표', '치수표 수동 입력 완료')
    replace_once(presentation.slides[9], 'OCR 원문 확인·수정 후 S/M/L별 총장·어깨·가슴 비교',
                 'S/M/L별 총장·어깨·가슴 수동 입력·비교; OCR은 후속')
    replace_once(presentation.slides[9], '합성 조정 실험', '합성 후처리 실험')
    replace_once(presentation.slides[9], '어깨 마스크·세로 길이·가슴 영역 각각 효과 검증',
                 '총장·몸통 폭 변형 연결; 실제 효과·복원 품질 검증 필요')


def add_comparison(presentation):
    slide = header(presentation, '11 / 이후 구현',
                   '판매 사이즈표와 두 입력 경로 연결',
                   '기준 옷 실측 또는 Ridge 예측값을 새 옷 S/M/L 치수와 비교')
    card(slide, .72, 2.1, 3.82, 3.65, '01  입력', '기준 옷이 있을 때', [
        '잘 맞는 옷의 총장·어깨너비·', '가슴단면을 직접 입력', '새 옷 − 기준 옷 차이 표시'])
    card(slide, 4.76, 2.1, 3.82, 3.65, '02  대체 경로', '기준 옷이 없을 때', [
        '성별·키·체중으로 Ridge 예측', '신체 치수의 예측 범위 표시', '옷 가슴단면과 여유 비교'])
    card(slide, 8.8, 2.1, 3.82, 3.65, '03  상품', '실제 판매 사이즈', [
        'S/M/L 등 라벨과 행 순서 보존', '총장·어깨너비·가슴단면 입력', '모호한 어깨 비교는 보류'])
    box(slide, .72, 6.06, 11.9, .54, PALE, PALE)
    text(slide, '완료: 직접 입력·검증·비교  |  남음: OCR 연결·적정 여유 기준·확정 추천',
         .98, 6.19, 11.4, .28, 13, INK)
    slide.notes_slide.notes_text_frame.text = (
        '근거: app/garment_compare.py, scripts/debug_size_compare.py, AGENTS.md. '
        '신체 가슴둘레/2는 옷 가슴단면과 비교하기 위한 값이며 사진 정면 폭이 아니다.'
    )


def add_pipeline(presentation):
    slide = header(presentation, '12 / 합성 흐름',
                   'FASHN 기본 합성 한 번 → S/M/L 후처리',
                   '입력 치수가 FASHN 내부 조건이 되는 방식은 아니며, 생성 결과의 옷 영역을 조정')
    steps = [
        ('1', '기준 사진·5점', '목·몸통 좌우·밑단 좌우 확인'),
        ('2', 'FASHN 기본 합성', '인물 + 상품 사진으로 한 번 생성'),
        ('3', '옷 마스크 확인', 'GrabCut 초안에 브러시·지우개 적용'),
        ('4', '목표선·변형', '총장 비율·가슴단면 비율 적용'),
        ('5', '결과 비교', 'S/M/L 갤러리·선택 결과 저장'),
    ]
    for index, (number, title, detail) in enumerate(steps):
        x = .72 + index*2.46
        box(slide, x, 2.2, 2.28, 2.55)
        text(slide, number, x+.18, 2.42, .38, .48, 22, TEAL, True)
        text(slide, title, x+.18, 3.02, 1.95, .72, 16, INK, True)
        text(slide, detail, x+.18, 3.83, 1.95, .68, 12, MUTED)
        if index < 4:
            text(slide, '→', x+2.29, 3.21, .18, .4, 16, TEAL, True)
    box(slide, .72, 5.16, 11.9, 1.26, PALE, PALE)
    text(slide, '현재 확인된 변화', 1, 5.35, 3.1, .4, 16, TEAL, True)
    text(slide, '선택한 판매 사이즈에 따라 총장 목표와 출력 이미지가 달라짐',
         4.0, 5.37, 8.25, .4, 15, INK)
    text(slide, '몸통 폭의 실제 효과는 미확인 · 어깨 변형은 보류',
         1, 5.85, 11.2, .35, 12, MUTED)
    slide.notes_slide.notes_text_frame.text = (
        '근거: app/fit_plan.py, app/garment_outline.py, app/size_synthesis.py, '
        'scripts/debug_size_preview.py, scripts/garment_mask_editor.py. '
        '로컬은 CPU 변형과 모의 FASHN·브라우저 연결을 검증했다. '
        '사용자는 Colab 실제 합성에서 총장 변화를 확인했지만 폭 변화는 확인하지 못했다.'
    )


def add_repair(presentation):
    slide = header(presentation, '13 / 실패와 개선 실험',
                   '짧아진 옷 아래 드러난 영역 복원',
                   '실제 합성에서 허리·바지·배경 연결이 부자연스러운 문제를 확인')
    card(slide, .72, 2.04, 5.72, 3.48, '이전 실험', 'Telea 복원', [
        '기존 옷을 지운 뒤 주변 색으로 채움',
        '밑단 아래로 피부·배경색이 삼각형처럼 번짐',
        '실제 사진에서 경계가 어색함'], ORANGE)
    card(slide, 6.65, 2.04, 5.97, 3.48, '현재 기본값', '원본 재사용 + 방향성 패치', [
        '보이는 원본만 조건부로 재사용',
        '밑단 아래·옆의 패치를 해당 방향으로 연장',
        '초록=재사용 / 주황=추정 부분'])
    box(slide, .72, 5.84, 11.9, .68, PALE, PALE)
    text(slide, '인공 줄무늬 사례: 복원 MAE 56.29 → 0.00  |  실제 사진 품질을 뜻하지 않음',
         .98, 6.0, 11.3, .36, 13, INK)
    slide.notes_slide.notes_text_frame.text = (
        '근거: app/garment_repair.py, docs/PLAN.md의 드러난 영역 복원 기록. '
        '수치는 220×260px 인공 줄무늬 바지 통제 사례의 3,240px 복원 영역에 한정된다. '
        '사용자의 실제 사진에서는 현재 패치 방식도 부자연스럽다고 관찰했다.'
    )


def add_status(presentation):
    slide = header(presentation, '14 / 검증 상태와 다음 실험',
                   '동작은 연결됐고, 자연스러운 착용 결과는 검증 중',
                   '소프트웨어 동작 검증과 실착 핏·복원 품질 검증을 구분')
    card(slide, .72, 2.06, 3.82, 3.82, '확인된 동작', '구현', [
        'Ridge 예측·치수표 직접 비교',
        '5점 목표선·FASHN 후처리 연결',
        '마스크 편집·사이즈 갤러리',
        '기본 합성 재사용'])
    card(slide, 4.76, 2.06, 3.82, 3.82, '현재 한계', '실제 사진', [
        '총장 변화 관찰',
        '몸통 폭 효과는 추가 확인 필요',
        '어깨 변형은 보류',
        '가려진 허리·바지 복원은 어색함'], ORANGE)
    card(slide, 8.8, 2.06, 3.82, 3.82, '다음 실험 제안', '부분 복원', [
        '드러난 부위만 생성형 복원 시험',
        'S/M/L에서 같은 복원 결과 재사용',
        'Colab T4 품질·시간 직접 측정',
        '실착 사진으로 경계·핏 비교'])
    text(slide, '생성형 인페인팅은 검토 중인 아이디어이며 현재 코드에는 연결하지 않았다.',
         .79, 6.2, 11.8, .33, 13, MUTED)
    slide.notes_slide.notes_text_frame.text = (
        '후속 후보: 드러난 영역에만 생성형 인페인팅을 시험한다. '
        '첫 복원 연산 시간은 추가되므로 T4에서 실제로 측정해야 한다. '
        '복원 모델이 원래 옷 밑에 가려진 실제 신체 모습을 알아내는 것은 아니다. '
        '외부 사용자 정확도, 추천 기준, OCR, 어깨/가슴 변형 효과도 미검증이다.'
    )


def main():
    presentation = Presentation(SOURCE)
    if len(presentation.slides) != 11:
        raise ValueError('원본 PPT의 장수가 바뀌어 수동 확인이 필요합니다.')
    update_old_status(presentation)
    add_comparison(presentation)
    add_pipeline(presentation)
    add_repair(presentation)
    add_status(presentation)
    presentation.save(OUTPUT)
    print(f'{len(presentation.slides)}장 저장: {OUTPUT}')


if __name__ == '__main__':
    main()
